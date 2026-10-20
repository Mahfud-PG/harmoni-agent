"""
Harmoni Agent - Document & File Parser Engine
Mengekstrak teks dan konten dari berbagai format file:
- Dokumen: PDF, DOCX
- Teks & Lirik/Data: TXT, MD, LRC, SRT, CSV, JSON, LOG, PY, HTML
- Gambar: JPG, JPEG, PNG, WEBP
- Audio: MP3, WAV, OGG, FLAC, M4A (ekstraksi metadata via ffprobe)
"""

import os
import sys
import json
import subprocess
from pathlib import Path


def extract_text_from_pdf(file_path: str, max_pages: int = 50) -> str:
    """Mengekstrak teks dari file PDF menggunakan PyPDF2."""
    import PyPDF2
    text_parts = []
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            total_pages = len(reader.pages)
            pages_to_read = min(total_pages, max_pages)
            for i in range(pages_to_read):
                p_text = reader.pages[i].extract_text() or ""
                p_text_clean = p_text.strip()
                if p_text_clean:
                    text_parts.append(f"--- Halaman {i+1} ---\n{p_text_clean}")
            
            res = "\n\n".join(text_parts)
            if not res.strip():
                return "[Dokumen PDF ini tidak memiliki lapisan teks, kemungkinan berupa pindaian gambar/scan]."
            if total_pages > max_pages:
                res += f"\n\n[Catatan: Hanya {max_pages} dari {total_pages} halaman pertama yang dibaca]."
            return res
    except Exception as e:
        return f"[Gagal membaca file PDF: {str(e)}]"


def extract_text_from_docx(file_path: str) -> str:
    """Mengekstrak teks paragraf dan tabel dari file Microsoft Word (.docx)."""
    import docx
    try:
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        tables_text = []
        for t_idx, table in enumerate(doc.tables, 1):
            rows_data = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(" | ".join(row_cells))
            if rows_data:
                tables_text.append(f"[Tabel {t_idx}]:\n" + "\n".join(rows_data))
                
        full_text = "\n\n".join(paragraphs)
        if tables_text:
            full_text += "\n\n--- TABEL DOKUMEN ---\n" + "\n\n".join(tables_text)
            
        return full_text or "[Dokumen Word (.docx) ini kosong]."
    except Exception as e:
        return f"[Gagal membaca file DOCX: {str(e)}]"


def extract_text_from_plain(file_path: str, max_chars: int = 50000) -> str:
    """Mengekstrak isi file berbasis teks murni (TXT, MD, LRC, SRT, CSV, JSON)."""
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252", "gb18030"]:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read(max_chars)
                if len(content) >= max_chars:
                    content += f"\n\n[Catatan: Teks dibatasi hingga {max_chars} karakter pertama]."
                return content
        except (UnicodeDecodeError, UnicodeError):
            continue
    return "[Gagal mengidentifikasi encoding teks file ini]."


def extract_audio_metadata(file_path: str) -> str:
    """Membaca metadata lagu/audio (durasi, title, artist, sample rate) menggunakan ffprobe."""
    try:
        cmd = [
            "ffprobe", "-v", "error",
            "-show_entries", "format=duration,bit_rate,format_name,size:format_tags=title,artist,album,genre,date",
            "-show_entries", "stream=codec_name,sample_rate,channels",
            "-of", "json",
            file_path
        ]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="ignore")
        if r.returncode == 0:
            data = json.loads(r.stdout)
            fmt = data.get("format", {})
            tags = fmt.get("tags", {})
            streams = data.get("streams", [{}])
            audio_stream = next((s for s in streams if s.get("codec_name")), {})
            
            durasi_detik = float(fmt.get("duration", 0))
            menit = int(durasi_detik // 60)
            detik = int(durasi_detik % 60)
            
            info_lines = [
                f"🎵 **Informasi File Audio:**",
                f"- Nama: {Path(file_path).name}",
                f"- Judul: {tags.get('title') or tags.get('TITLE') or '-'}",
                f"- Artis: {tags.get('artist') or tags.get('ARTIST') or '-'}",
                f"- Album: {tags.get('album') or tags.get('ALBUM') or '-'}",
                f"- Durasi: {menit} menit {detik} detik ({durasi_detik:.1f} detik)",
                f"- Format/Codec: {audio_stream.get('codec_name', fmt.get('format_name', '-')).upper()}",
                f"- Sample Rate: {audio_stream.get('sample_rate', '-')} Hz",
                f"- Channel: {'Stereo (2)' if audio_stream.get('channels') == 2 else 'Mono (1)'}",
                f"- Bitrate: {int(fmt.get('bit_rate', 0)) // 1000 if fmt.get('bit_rate') else '-'} kbps"
            ]
            return "\n".join(info_lines)
    except Exception as e:
        pass
    return f"[File audio {Path(file_path).name} terdeteksi, durasi/metadata tidak dapat diekstrak]."


def extract_image_info(file_path: str) -> str:
    """Membaca informasi dimensi dan format gambar."""
    try:
        from PIL import Image
        with Image.open(file_path) as img:
            w, h = img.size
            fmt = img.format
            mode = img.mode
            return (
                f"🖼️ **Informasi File Gambar/Foto:**\n"
                f"- File: {Path(file_path).name}\n"
                f"- Format: {fmt}\n"
                f"- Dimensi: {w} x {h} piksel\n"
                f"- Mode Warna: {mode}\n"
                f"- Catatan: Gambar telah siap dianalisis secara visual oleh model AI."
            )
    except Exception as e:
        return f"[File gambar {Path(file_path).name}, error baca header: {str(e)}]"


def parse_document(file_path: str) -> dict:
    """
    Dispatcher utama untuk membaca dan mengekstrak isi file berdasarkan ekstensinya.
    Mengembalikan dict:
    {
        "status": "success" | "error",
        "file_name": str,
        "file_ext": str,
        "category": "document" | "text" | "audio" | "image" | "unknown",
        "file_size_kb": float,
        "extracted_text": str,
        "summary_prompt": str
    }
    """
    path_obj = Path(file_path)
    if not path_obj.exists():
        return {
            "status": "error",
            "file_name": path_obj.name,
            "error": "File tidak ditemukan di komputer."
        }

    ext = path_obj.suffix.lower()
    size_kb = round(path_obj.stat().st_size / 1024, 2)
    
    extracted = ""
    category = "unknown"
    
    # 1. Dokumen PDF
    if ext == ".pdf":
        category = "document"
        extracted = extract_text_from_pdf(str(path_obj))
        
    # 2. Dokumen Word DOCX
    elif ext in [".docx", ".doc"]:
        category = "document"
        if ext == ".docx":
            extracted = extract_text_from_docx(str(path_obj))
        else:
            extracted = "[Format .doc lama tidak didukung langsung, mohon simpan sebagai .docx atau .pdf]."
            
    # 3. Teks Murni & Lirik
    elif ext in [".txt", ".md", ".lrc", ".srt", ".csv", ".json", ".xml", ".html", ".log", ".py", ".js", ".css"]:
        category = "text"
        extracted = extract_text_from_plain(str(path_obj))
        
    # 4. Audio
    elif ext in [".mp3", ".wav", ".ogg", ".flac", ".m4a", ".aac", ".wma"]:
        category = "audio"
        extracted = extract_audio_metadata(str(path_obj))
        
    # 5. Gambar
    elif ext in [".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"]:
        category = "image"
        extracted = extract_image_info(str(path_obj))
        
    else:
        category = "unknown"
        extracted = f"[Format file '{ext}' belum didukung untuk ekstraksi teks langsung]."

    # Siapkan prompt rangkuman terstruktur untuk AI
    summary_prompt = (
        f"\n\n--- LAMPIRAN DOKUMEN PENGGUNA ---\n"
        f"Nama File: {path_obj.name}\n"
        f"Kategori: {category.upper()}\n"
        f"Ukuran: {size_kb} KB\n\n"
        f"ISI/DATA FILE:\n{extracted}\n"
        f"--- AKHIR DOKUMEN ---\n\n"
    )

    return {
        "status": "success",
        "file_name": path_obj.name,
        "file_ext": ext,
        "category": category,
        "file_size_kb": size_kb,
        "extracted_text": extracted,
        "summary_prompt": summary_prompt,
        "file_path": str(path_obj)
    }


if __name__ == "__main__":
    print("Testing document_parser...")
    # Test on a dummy text file
    dummy_txt = Path(__file__).resolve().parent.parent / "File" / "Temp" / "test_sample.txt"
    dummy_txt.parent.mkdir(parents=True, exist_ok=True)
    dummy_txt.write_text("Judul: Perjanjian Rekaman Studio Harmoni\nBiaya: Rp 750.000\nTanggal: 26 September 2026", encoding="utf-8")
    
    res = parse_document(str(dummy_txt))
    print("Parsed Text File Result:")
    print("Category:", res["category"])
    print("Extracted:", res["extracted_text"])
    dummy_txt.unlink()
    print("Test OK!")
